#!/usr/bin/env python3
"""Build the talk deck (deck.pptx) from deck.md, with script.md in the notes.

The deck is black text on a white ground in one typeface. The only things on
a slide besides text boxes are the author's pictures and one table, kept from
the author's deck of 2026-09-25: no boxes, bars, cards, icons, colour on the
text, or transitions. Every run is given an explicit black colour and every
slide an explicit white background, so a venue template cannot tint anything.

deck.md holds one block per slide:

    ## Slide 3 — the dilemma
    #### FOUNDATIONAL PROBLEM          (small label above the headline)
    ### The Ontological Reducibility Dilemma
    **Hypothesis A: Total Reducibility**
    Dyadic Collapse: All systems reduce to pairs.

    **Hypothesis B: Genuine Irreducibility**
    ![caption](media/file.png)        (picture in the left column)
    | a | b |                         (table rows; the first is the header)

A blank line starts a new group; consecutive lines sit tight. A leading
``**...**`` is set bold. script.md uses the same ``## Slide k`` headings;
each section becomes that slide's speaker notes, with HTML comments removed.
Run with a Python that has python-pptx:

    /usr/bin/python3 talk/build_deck.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
DECK_MD = HERE / "deck.md"
SCRIPT_MD = HERE / "script.md"
OUT = HERE / "deck.pptx"

BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SANS = "Arial"
MATH = "Cambria Math"  # ships with Office; carries the symbols Arial lacks
MATH_CHARS = set("∈⊕")
WIDTH, HEIGHT = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.9)
PICTURE_W = Inches(3.4)
GUTTER = Inches(0.5)

SIZES = {"label": 14, "headline": 34, "body": 22, "url": 12, "caption": 16, "table": 16,
         "title": 44, "subtitle": 22}

SLIDE_RE = re.compile(r"^## Slide (\d+)\b.*$", re.M)
PICTURE_RE = re.compile(r"^!\[(.*?)\]\((.+?)\)$")
BOLD_RE = re.compile(r"^\*\*(.+?)\*\*(.*)$")


def split_slides(text: str) -> dict[int, str]:
    parts = {}
    marks = list(SLIDE_RE.finditer(text))
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        parts[int(m.group(1))] = text[m.end():end].strip("\n")
    return parts


def notes_text(section: str) -> str:
    """Speaker notes: the script section without HTML comments."""
    text = re.sub(r"<!--.*?-->\n?", "", section, flags=re.S)
    text = re.sub(r"^---$", "", text, flags=re.M)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def parse_slide(block: str) -> dict:
    """Split one deck.md block into label, headline, body lines, picture, table."""
    slide = {"label": "", "headline": "", "body": [], "picture": None, "table": []}
    gap = False
    for line in block.splitlines():
        if line.startswith("<!--"):
            continue
        if not line.strip():
            gap = True
            continue
        if line.startswith("#### ") and not slide["label"]:
            slide["label"] = line[5:].strip()
        elif line.startswith("### ") and not slide["headline"]:
            slide["headline"] = line[4:].strip()
        elif PICTURE_RE.match(line):
            caption, path = PICTURE_RE.match(line).groups()
            slide["picture"] = (caption, HERE / path)
        elif line.startswith("|"):
            slide["table"].append([c.strip() for c in line.strip().strip("|").split("|")])
        else:
            slide["body"].append((gap and bool(slide["body"]), line.strip()))
            gap = False
    return slide


def white_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def style_run(run, style: str, bold: bool = False, face: str = SANS) -> None:
    font = run.font
    font.name = face
    font.size = Pt(SIZES[style])
    font.bold = bold or style in ("headline", "title")
    font.italic = False
    font.color.rgb = BLACK


def add_runs(p, text: str, style: str) -> None:
    if text.startswith(("http://", "https://")):
        style = "url"  # a bare URL cannot wrap at a space
    m = BOLD_RE.match(text)
    pieces = [(m.group(1), True), (m.group(2), False)] if m else [(text, False)]
    for piece, bold in pieces:
        # Split out the symbols Arial cannot draw and set them in the math face.
        for seg in re.findall(r"[∈⊕]+|[^∈⊕]+", piece):
            face = MATH if seg[0] in MATH_CHARS else SANS
            style_run(p.add_run(), style, bold, face)
            p.runs[-1].text = seg


def add_box(slide, left, top, width, height, lines, style, anchor=MSO_ANCHOR.TOP,
            align=PP_ALIGN.LEFT):
    """lines: [(starts_group, text)]."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(0))
    for k, (starts_group, text) in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0 if k == 0 else (16 if starts_group else 4))
        add_runs(p, text, style)
    return box


def plain_table(slide, rows, left, top, width) -> None:
    """A table with black text, no fill, and thin black rules."""
    shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width,
                                   Inches(0.55) * len(rows))
    table = shape.table
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tbl_pr.set(attr, "0")
    style_id = tbl_pr.find(qn("a:tableStyleId"))
    if style_id is not None:
        style_id.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.background()
            cell.margin_left = cell.margin_right = Inches(0.08)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            style_run(p.add_run(), "table", bold=(r == 0))
            p.runs[0].text = text
            tc_pr = cell._tc.get_or_add_tcPr()
            for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = tc_pr.makeelement(qn(edge), {"w": "9525"})
                solid = ln.makeelement(qn("a:solidFill"), {})
                solid.append(solid.makeelement(qn("a:srgbClr"), {"val": "000000"}))
                ln.append(solid)
                tc_pr.append(ln)


def add_picture(slide, path: Path, caption: str, left, top, max_h) -> None:
    w, h = Image.open(path).size
    width = PICTURE_W
    height = int(width * h / w)
    if height > max_h:
        height, width = max_h, int(max_h * w / h)
    slide.shapes.add_picture(str(path), left, top, width, height)
    if caption:
        add_box(slide, left, top + height + Inches(0.15), PICTURE_W, Inches(0.6),
                [(False, caption)], "caption")


def build() -> Presentation:
    deck = split_slides(DECK_MD.read_text(encoding="utf-8"))
    script = split_slides(SCRIPT_MD.read_text(encoding="utf-8")) if SCRIPT_MD.exists() else {}
    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDTH, HEIGHT
    blank = prs.slide_layouts[6]
    inner_w = WIDTH - 2 * MARGIN
    for n in sorted(deck):
        s = parse_slide(deck[n])
        slide = prs.slides.add_slide(blank)
        white_background(slide)
        if n == 1:
            if s["label"]:
                add_box(slide, MARGIN, Inches(1.6), inner_w, Inches(0.4),
                        [(False, s["label"])], "label", anchor=MSO_ANCHOR.BOTTOM)
            add_box(slide, MARGIN, Inches(2.1), inner_w, Inches(1.2),
                    [(False, s["headline"])], "title", anchor=MSO_ANCHOR.BOTTOM)
            add_box(slide, MARGIN, Inches(3.6), inner_w, Inches(3.0), s["body"], "subtitle")
        else:
            if s["label"]:
                add_box(slide, MARGIN, Inches(0.55), inner_w, Inches(0.3),
                        [(False, s["label"])], "label")
            add_box(slide, MARGIN, Inches(0.9), inner_w, Inches(1.1),
                    [(False, s["headline"])], "headline")
            top = Inches(2.2)
            left, width = MARGIN, inner_w
            if s["picture"]:
                caption, path = s["picture"]
                add_picture(slide, path, caption, MARGIN, top, HEIGHT - top - Inches(1.2))
                left = MARGIN + PICTURE_W + GUTTER
                width = WIDTH - MARGIN - left
            if s["body"]:
                add_box(slide, left, top, width, HEIGHT - top - Inches(0.6), s["body"], "body")
            if s["table"]:
                plain_table(slide, s["table"], MARGIN, top + Inches(0.3), inner_w)
        slide.notes_slide.notes_text_frame.text = notes_text(script.get(n, ""))
    return prs


def main() -> int:
    prs = build()
    prs.save(OUT)
    print(f"wrote {OUT} ({len(prs.slides)} slides)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
