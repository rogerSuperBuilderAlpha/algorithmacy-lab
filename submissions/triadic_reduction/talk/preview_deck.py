#!/usr/bin/env python3
"""Draw each slide of deck.pptx to a PNG and report text that overflows.

There is no PowerPoint on the build machine, so this is an approximation:
it lays out every text box from the .pptx with the real Arial and Courier
New metrics, wraps words to the box width, and flags any box whose text runs
past its bottom edge. Good enough to catch overflow and crowding; the author
still looks at the deck in PowerPoint or Keynote before presenting.

    /usr/bin/python3 talk/preview_deck.py OUT_DIR
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
FONT_DIR = Path("/System/Library/Fonts/Supplemental")
FACES = {
    ("Arial", False, False): "Arial.ttf", ("Arial", True, False): "Arial Bold.ttf",
    ("Arial", False, True): "Arial Italic.ttf", ("Arial", True, True): "Arial Bold Italic.ttf",
    ("Courier New", False, False): "Courier New.ttf",
    ("Cambria Math", False, False): "STIXTwoMath.otf",
}
DPI = 96  # pixels per inch in the preview


def px(emu) -> int:
    return int(Emu(emu).inches * DPI)


def font_for(run) -> ImageFont.FreeTypeFont:
    key = (run.font.name or "Arial", bool(run.font.bold), bool(run.font.italic))
    path = FONT_DIR / FACES.get(key, FACES.get((key[0], False, False), "Arial.ttf"))
    size_pt = run.font.size.pt if run.font.size else 18
    return ImageFont.truetype(str(path), int(size_pt * DPI / 72))


def wrap(draw, text, font, width) -> list[str]:
    lines, line = [], ""
    for word in text.split(" "):
        trial = f"{line} {word}".strip()
        if draw.textlength(trial, font=font) <= width or not line:
            line = trial
        else:
            lines.append(line)
            line = word
    lines.append(line)
    return lines


def render(prs, out: Path) -> list[str]:
    out.mkdir(parents=True, exist_ok=True)
    w, h = px(prs.slide_width), px(prs.slide_height)
    warnings = []
    for i, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (w, h), "white")
        draw = ImageDraw.Draw(img)
        for shape in slide.shapes:
            x, y = px(shape.left), px(shape.top)
            bw, bh = px(shape.width), px(shape.height)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pic = Image.open(io.BytesIO(shape.image.blob)).convert("RGBA")
                img.paste(pic.resize((bw, bh)), (x, y), pic.resize((bw, bh)))
                continue
            if shape.has_table:
                cy = y
                for row in shape.table.rows:
                    cx, rh = x, px(row.height)
                    for col, cell in zip(shape.table.columns, row.cells):
                        cw = px(col.width)
                        run = cell.text_frame.paragraphs[0].runs[0]
                        font = font_for(run)
                        lines = wrap(draw, cell.text_frame.text, font, cw - 16)
                        rh = max(rh, len(lines) * int(font.size * 1.2) + 12)
                    for col, cell in zip(shape.table.columns, row.cells):
                        cw = px(col.width)
                        draw.rectangle([cx, cy, cx + cw, cy + rh], outline="black")
                        font = font_for(cell.text_frame.paragraphs[0].runs[0])
                        ty = cy + 6
                        for line in wrap(draw, cell.text_frame.text, font, cw - 16):
                            draw.text((cx + 8, ty), line, font=font, fill="black")
                            ty += int(font.size * 1.2)
                        cx += cw
                    cy += rh
                continue
            paras = []
            for p in shape.text_frame.paragraphs:
                if not p.runs:
                    continue
                run = max(p.runs, key=lambda r: len(r.text))
                font = font_for(run)
                gap = int(p.space_before.pt * DPI / 72) if p.space_before else 0
                paras.append((gap, font, wrap(draw, p.text, font, bw)))
            height = sum(g + len(ls) * int(f.size * 1.2) for g, f, ls in paras)
            anchor_bottom = shape.text_frame.vertical_anchor == 4  # MSO_ANCHOR.BOTTOM
            cy = y + (bh - height if anchor_bottom and height < bh else 0)
            for gap, font, lines in paras:
                cy += gap
                for line in lines:
                    draw.text((x, cy), line, font=font, fill="black")
                    cy += int(font.size * 1.2)
            if height > bh:
                warnings.append(f"slide {i}: text overflows its box by {height - bh}px")
        img.save(out / f"slide-{i:02d}.png")
    return warnings


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: preview_deck.py OUT_DIR   (keep previews out of the repo)")
        return 2
    out = Path(sys.argv[1])
    prs = Presentation(HERE / "deck.pptx")
    warnings = render(prs, out)
    for wmsg in warnings:
        print("OVERFLOW", wmsg)
    print(f"wrote {len(prs.slides)} previews to {out}; {len(warnings)} overflow warnings")
    return 1 if warnings else 0


if __name__ == "__main__":
    sys.exit(main())
